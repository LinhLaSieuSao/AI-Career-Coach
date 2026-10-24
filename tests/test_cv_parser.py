from io import BytesIO

import docx
import pytest

from modules import parse_document
from modules.cv_parser import extract_text, extract_text_from_pdf, extract_text_from_docx
from document_fixtures import pdf_bytes, docx_bytes

def _minimal_pdf_bytes(text: str) -> bytes:
    escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({escaped_text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream"),
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    )
    pdf.extend(trailer.encode("ascii"))
    return bytes(pdf)

def _docx_bytes(document: docx.Document) -> BytesIO:
    source = BytesIO()
    document.save(source)
    source.seek(0)
    return source

def test_unsupported_format_raises_value_error():
    with pytest.raises(ValueError, match="Định dạng tệp không được hỗ trợ"):
        extract_text("resume.txt")

def test_missing_pdf_path_returns_empty_string():
    assert extract_text("khong_ton_tai.pdf") == ""

def test_extract_valid_pdf_from_bytesio():
    source = BytesIO(_minimal_pdf_bytes("Python developer"))
    source.seek(0, 2)  # Parser must rewind an already-consumed upload.
    assert "Python developer" in extract_text(source, filename="resume.pdf")

def test_extract_docx_from_bytesio():
    document = docx.Document()
    document.add_paragraph("Python developer")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "SQL and data engineering"
    extracted = extract_text(_docx_bytes(document), filename="resume.docx")
    assert "Python developer" in extracted
    assert "SQL and data engineering" in extracted


@pytest.mark.parametrize("source", [None, "", "missing.pdf"])
def test_structured_parser_no_file(source):
    result = parse_document(source)
    assert result["text"] == ""
    assert result["error_code"] == "NO_FILE"
    assert result["message"]


@pytest.mark.parametrize("name,content,code", [
    ("valid.PDF", pdf_bytes("Python"), None),
    ("valid.DOCX", docx_bytes("Python"), None),
    ("empty.pdf", b"", "EMPTY"),
    ("empty.docx", b"", "EMPTY"),
    ("blank.pdf", pdf_bytes(), "EMPTY"),
    ("drawing.pdf", pdf_bytes(vector=True), "EMPTY"),
    ("blank.docx", docx_bytes(), "EMPTY"),
    ("broken.pdf", b"%PDF-1.4 broken", "UNREADABLE"),
    ("broken.docx", b"not a zip file", "UNREADABLE"),
    ("image.pdf", pdf_bytes(image=True), "SCANNED_PDF"),
    ("other.txt", b"Python", "UNSUPPORTED_FORMAT"),
    ("no_extension", pdf_bytes("Python"), "UNSUPPORTED_FORMAT"),
    ("fake.pdf", docx_bytes("Python"), "UNREADABLE"),
    ("fake.docx", pdf_bytes("Python"), "UNREADABLE"),
], ids=["valid-pdf", "valid-docx", "zero-pdf", "zero-docx", "blank-pdf", "vector-pdf", "blank-docx", "broken-pdf", "broken-docx", "image-pdf", "text-file", "no-suffix", "docx-as-pdf", "pdf-as-docx"])
def test_structured_parser_paths_and_consumed_streams(tmp_path, name, content, code):
    path = tmp_path / name
    path.write_bytes(content)
    stream = BytesIO(content)
    stream.name = name
    stream.seek(0, 2)
    results = [parse_document(path), parse_document(stream)]
    assert results[0] == results[1]
    for result in results:
        assert result["error_code"] == code
        assert result["text"] == ("Python" if code is None else "")
        assert bool(result["message"]) == (code is not None)
    assert not stream.closed


def test_explicit_filename_overrides_stream_name():
    stream = BytesIO(pdf_bytes("SQL"))
    stream.name = "wrong.txt"
    assert parse_document(stream, filename="CV.PDF")["text"] == "SQL"
    assert parse_document(stream)["error_code"] == "UNSUPPORTED_FORMAT"


def test_unusable_stream_is_unreadable():
    stream = BytesIO(b"data")
    stream.close()
    assert parse_document(stream, "cv.pdf")["error_code"] == "UNREADABLE"


def test_docx_paragraph_table_nested_table_and_merged_cells_in_order():
    document = docx.Document()
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "Merged table cell"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested SQL"
    cell.add_paragraph("After nested table")
    document.add_paragraph("Last Python paragraph")
    source = _docx_bytes(document)
    result = parse_document(source, "nested.docx")
    assert result["error_code"] is None
    assert [line for line in result["text"].splitlines() if line] == [
        "First paragraph", "Merged table cell", "Nested SQL", "After nested table", "Last Python paragraph"
    ]
    assert extract_text(source, "nested.docx") == result["text"]
    assert extract_text_from_docx(source) == result["text"]


def test_legacy_pdf_wrapper_and_structured_parser_share_reader():
    source = BytesIO(pdf_bytes("Python"))
    assert extract_text_from_pdf(source) == parse_document(source, "cv.pdf")["text"]
    assert extract_text_from_pdf(BytesIO(b"broken")) == ""
    assert extract_text_from_docx(BytesIO(b"broken")) == ""
