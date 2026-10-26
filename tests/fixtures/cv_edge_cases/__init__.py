"""Create synthetic CV edge-case files without storing personal documents."""

from __future__ import annotations

from pathlib import Path

import docx

from document_fixtures import pdf_bytes


def _unicode_pdf_bytes(lines: list[tuple[int, int, str]]) -> bytes:
    """Build a small text PDF with a ToUnicode map for deterministic fixtures."""
    characters = sorted({char for _, _, text in lines for char in text})
    content_parts = ["BT /F0 11 Tf"]
    for x, y, text in lines:
        encoded = text.encode("utf-16-be").hex().upper()
        content_parts.append(f"1 0 0 1 {x} {y} Tm <{encoded}> Tj")
    content_parts.append("ET")
    content = "\n".join(content_parts).encode("ascii")
    mappings = "\n".join(
        f"<{ord(char):04X}> <{char.encode('utf-16-be').hex().upper()}>"
        for char in characters
    )
    cmap = (
        "/CIDInit /ProcSet findresource begin\n"
        "12 dict begin\nbegincmap\n"
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
        "/CMapName /AICC-Identity-UCS def\n/CMapType 2 def\n"
        "1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n"
        f"{len(characters)} beginbfchar\n{mappings}\nendbfchar\n"
        "endcmap\nCMapName currentdict /CMap defineresource pop\nend\nend"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F0 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type0 /BaseFont /AICCIdentity /Encoding /Identity-H /DescendantFonts [6 0 R] /ToUnicode 7 0 R >>",
        b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /AICCIdentity /CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> /FontDescriptor 8 0 R /DW 1000 >>",
        b"<< /Length " + str(len(cmap)).encode() + b" >>\nstream\n" + cmap + b"\nendstream",
        b"<< /Type /FontDescriptor /FontName /AICCIdentity /Flags 4 /FontBBox [0 -200 1000 900] /ItalicAngle 0 /Ascent 800 /Descent -200 /CapHeight 700 /StemV 80 >>",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(result)
    result.extend(f"xref\n0 {len(objects)+1}\n".encode("ascii"))
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    result.extend(
        f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii")
    )
    return bytes(result)


def create_cv_edge_cases(directory: str | Path) -> dict[str, Path]:
    """Write synthetic PDF/DOCX cases to a temporary directory for parser/UI QA."""
    root = Path(directory)
    root.mkdir(parents=True, exist_ok=True)
    files = {
        "valid_vietnamese.pdf": _unicode_pdf_bytes([(48, 730, "Nguyễn An · Chuyên viên phân tích dữ liệu")]),
        "two_column_resume.pdf": _unicode_pdf_bytes([
            (48, 730, "KỸ NĂNG · Python, SQL"),
            (320, 730, "KINH NGHIỆM · Báo cáo dữ liệu"),
        ]),
        "scanned_image.pdf": pdf_bytes(image=True),
        "empty.pdf": b"",
        "wrong_magic.pdf": b"MZ\x90\x00synthetic executable payload",
    }
    doc = docx.Document()
    doc.add_paragraph("Nguyễn An · Kỹ sư phần mềm")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Kinh nghiệm với dữ liệu có dấu"
    output = root / "docx_table_resume.docx"
    doc.save(output)
    result = {name: root / name for name in files}
    for name, content in files.items():
        result[name].write_bytes(content)
    result[output.name] = output
    result["empty.docx"] = root / "empty.docx"
    result["empty.docx"].write_bytes(b"")
    return result
